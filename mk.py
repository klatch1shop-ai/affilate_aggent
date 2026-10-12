# -*- coding: utf-8 -*-
import json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open('/tmp/rd_sel.json'))
взяті, відкинуті = d['взяті'], d['відкинуті']
ваги = [v['вага'] for v in взяті if v['вага']]
print('відібрано:', len(взяті), '· з вагою:', len(ваги),
      '· макс вага:', max(ваги) if ваги else '—')
print('найважчі серед відібраних:')
for v in sorted([x for x in взяті if x['вага']], key=lambda x: -x['вага'])[:5]:
    print(f"   {v['вага']:5} кг  {v['cat'][:26]:26} {v['name'][:48]}")

ARIAL = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864')
hdr_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ── Лист 1: правила ───────────────────────────────────────────────
ws = wb.active
ws.title = 'Правила'
rules = [
    ['ROZETKA Delivery — заявка на активацію', ''],
    ['', ''],
    ['Що це', 'Доставка в магазини ROZETKA. Продавець привозить посилку в магазин '
              'ROZETKA або у відділення Meest, покупець забирає в обраному магазині.'],
    ['Як вмикається', 'API для вмикання по окремих товарах НЕМАЄ. Єдиний шлях — '
                      'звернення в підтримку, тема «Видача посилок у ROZETKA», '
                      'з Excel, де вказані ID товарів на сайті ROZETKA.'],
    ['Що в цьому файлі', 'Стовпець «ID товару на ROZETKA» — це і є номер картки, '
                         'який потрібен підтримці.'],
    ['', ''],
    ['ОБМЕЖЕННЯ ПОСИЛКИ', ''],
    ['Вага', '0,1 – 30 кг'],
    ['Кожна сторона', '1 – 120 см'],
    ['Об’ємна вага', 'Д × Ш × В / 4000'],
    ['Пакування', 'Коробка з ребрами жорсткості; наліпки «Крихке»/«Верх» 10×10 см '
                  'на двох найбільших сторонах; накладна на найбільшій стороні. '
                  'Невідповідні відправлення не приймають.'],
    ['Вартість', 'Для продавця від 35 грн; для покупця безкоштовно; '
                 'повернення безкоштовне. Накладений платіж збирає ROZETKA Delivery.'],
    ['', ''],
    ['ЧЕСНО ПРО МЕЖІ ЦЬОГО ВІДБОРУ', ''],
    ['Вага відома не всюди', f'У фіді TOPTUL вага є лише у 451 з 5988 позицій '
                             f'({451/5988*100:.1f}%). Серед відібраних вага відома '
                             f'у {len(ваги)}. Решту відібрано за категорією, а не за вагою.'],
    ['Габарити невідомі', 'Розмірів у фіді немає взагалі. Обмеження 120 см '
                          'не перевірено — лише виключено явно масивні назви.'],
    ['Де перевірити напевно', 'Придатність конкретного товару видно в кабінеті: '
                              'модуль ROZETKA Delivery → Товари. Цей файл — '
                              'пропозиція до перевірки, а не підтверджений перелік.'],
]
for r in rules:
    ws.append(r)
ws['A1'].font = Font(name=ARIAL, bold=True, size=14)
for row in (7, 14):
    ws.cell(row=row, column=1).font = Font(name=ARIAL, bold=True, size=11)
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 95
for row in ws.iter_rows():
    for c in row:
        c.font = c.font or Font(name=ARIAL)
        if c.font.name != ARIAL:
            c.font = Font(name=ARIAL, bold=c.font.bold, size=c.font.size)
        c.alignment = Alignment(wrap_text=True, vertical='top')

# ── Листи з товарами ──────────────────────────────────────────────
COLS = [('ID товару на ROZETKA', 22), ('Артикул', 18), ('Назва', 62),
        ('Категорія', 28), ('Ціна, грн', 11), ('Вага, кг', 10),
        ('У наявності', 12), ('Посилання', 40)]

def лист(назва, рядки, підпис):
    w = wb.create_sheet(назва)
    w.append([підпис])
    w['A1'].font = Font(name=ARIAL, bold=True, size=11)
    w.append([])
    w.append([c[0] for c in COLS])
    for i, (_, width) in enumerate(COLS, 1):
        w.column_dimensions[get_column_letter(i)].width = width
        cell = w.cell(row=3, column=i)
        cell.fill, cell.font, cell.border = hdr_fill, hdr_font, box
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for v in рядки:
        w.append([v['rz_item_id'], v['article'], v['name'], v['cat'],
                  v['price'], v['вага'] if v['вага'] else '',
                  'так' if v['avail'] else 'ні', v['url']])
    for row in w.iter_rows(min_row=4):
        for c in row:
            c.font = Font(name=ARIAL, size=10)
            c.border = box
            c.alignment = Alignment(vertical='top',
                                    wrap_text=(c.column == 3))
    w.freeze_panes = 'A4'
    w.auto_filter.ref = f'A3:{get_column_letter(len(COLS))}{3 + len(рядки)}'
    return w

ручний = [v for v in взяті if v['група'] == 'ручний інструмент']
вимір = [v for v in взяті if v['група'] != 'ручний інструмент']
ручний.sort(key=lambda v: (v['cat'] or '', v['name'] or ''))
вимір.sort(key=lambda v: (v['cat'] or '', v['name'] or ''))
відкинуті.sort(key=lambda v: (v['чому'], v['cat'] or ''))

лист('Ручний інструмент', ручний,
     f'TOPTUL · ручний інструмент — {len(ручний)} позицій. Безумовно придатні: '
     f'ключі, головки, викрутки, біти, кліщі, знімачі, молотки.')
лист('Вимірювальні', вимір,
     f'TOPTUL · вимірювальні прилади й дрібна електроніка — {len(вимір)} позицій. '
     f'Невеликі, у заводських коробках.')

w = лист('Відкинуті', відкинуті,
         f'Не пропоную вмикати — {len(відкинуті)} позицій. Причина в останньому стовпці.')
w.cell(row=3, column=len(COLS) + 1, value='Чому відкинуто').fill = hdr_fill
w.cell(row=3, column=len(COLS) + 1).font = hdr_font
w.cell(row=3, column=len(COLS) + 1).border = box
w.column_dimensions[get_column_letter(len(COLS) + 1)].width = 34
for i, v in enumerate(відкинуті, 4):
    c = w.cell(row=i, column=len(COLS) + 1, value=v['чому'])
    c.font = Font(name=ARIAL, size=10); c.border = box

# ── Лист: інші постачальники ──────────────────────────────────────
ws2 = wb.create_sheet('Інші постачальники')
GOODS = json.load(open('/tmp/rz_goods.json'))
назви = {51949: 'Carvol (автоелектроніка)', 54788: 'NOIRE', 55861: 'dropoffice (будматеріали)'}
ws2.append(['Огляд інших прайсів — чи варто вмикати і їх'])
ws2['A1'].font = Font(name=ARIAL, bold=True, size=12)
ws2.append([])
ws2.append(['Постачальник', 'Карток', 'Найбільші категорії', 'Оцінка'])
for i in range(1, 5):
    c = ws2.cell(row=3, column=i); c.fill, c.font, c.border = hdr_fill, hdr_font, box
оцінки = {
 51949: 'Радше так: магнітоли, рамки, адаптери — компактні. Потребує окремої '
        'перевірки габаритів магнітол (екрани до 10").',
 54788: 'Радше так: товар дрібний. Але це 18+ — спершу спитати підтримку, '
        'чи приймають такі посилки у видачу в магазинах.',
 55861: 'Радше ні: панелі 700×770 мм і рулони плівки 10 м — близько до межі '
        '120 см і незручні в пакуванні.',
}
for src, підпис in назви.items():
    g = [x for x in GOODS if x['src'] == src]
    top = ', '.join(f'{k} ({v})' for k, v in Counter(x['cat'] for x in g).most_common(4))
    ws2.append([підпис, len(g), top, оцінки[src]])
for col, width in zip('ABCD', (28, 10, 60, 62)):
    ws2.column_dimensions[col].width = width
for row in ws2.iter_rows(min_row=4):
    for c in row:
        c.font = Font(name=ARIAL, size=10); c.border = box
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws2.append([])
ws2.append(['Пропозиція', 'Почати з TOPTUL: ручний інструмент — найбезпечніший '
            'випадок (легкий, компактний, у коробках). Коли ROZETKA Delivery '
            'запрацює, додати Carvol і NOIRE окремими заявками.'])
ws2.cell(row=ws2.max_row, column=1).font = Font(name=ARIAL, bold=True, size=10)
ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(wrap_text=True, vertical='top')

wb.save('/tmp/rozetka_delivery_toptul.xlsx')
print('\nзаписано /tmp/rozetka_delivery_toptul.xlsx')
print('листи:', wb.sheetnames)
