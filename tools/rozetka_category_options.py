#!/usr/bin/env python3
"""Характеристики категорій Rozetka → data/rozetka_category_options.json.

Rozetka віддає їх сама: `/market-categories/download-options?category_id=X`
повертає XLS із переліком характеристик категорії та допустимими
значеннями. Токен кабінету в нас є, тож завантаження нічого не коштує і
не потребує участі власника.

Потрібно для фіду: без обовʼязкових характеристик картка не пройде
модерацію, а дізнатись їх інакше можна лише очима в кабінеті.

    venv/bin/python tools/rozetka_category_options.py --категорії 4626670,4627031
    venv/bin/python tools/rozetka_category_options.py --з-каталогу-юк
"""
import argparse
import collections
import json
import os
import sys
import time

import requests

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
from agents.orders import rozetka_order_agent as RZ   # noqa: E402

API = 'https://api-seller.rozetka.com.ua'
OUT = os.path.join(BASE, 'data', 'rozetka_category_options.json')
ТЕКА = os.path.join(BASE, 'data', 'rozetka_options_xls')


def завантажити(cid, h):
    os.makedirs(ТЕКА, exist_ok=True)
    шлях = os.path.join(ТЕКА, f'{cid}.xls')
    if os.path.exists(шлях) and os.path.getsize(шлях) > 1000:
        return шлях
    r = requests.get(f'{API}/market-categories/download-options',
                     headers=h, params={'category_id': cid}, timeout=90)
    if r.status_code != 200 or len(r.content) < 1000:
        return None
    with open(шлях, 'wb') as f:
        f.write(r.content)
    return шлях


def розібрати(шлях):
    """{id параметра: {назва, тип, одиниця, значення[]}} з файла Rozetka.

    Будова файла: перший рядок — назва категорії, ДРУГИЙ — заголовок
    (ID параметра · Назва параметра · Тип параметра · Тип фільтра ·
    Одиниця вимірювання · ID значення), далі по рядку на кожне ДОПУСТИМЕ
    ЗНАЧЕННЯ. Тобто один параметр займає стільки рядків, скільки в нього
    значень, і читати перший рядок як заголовок не можна — так я спершу
    і зробив, отримавши «1 характеристику» замість трьох десятків.

    Rozetka віддає XLSX, але називає файл .xls, тому читаємо з памʼяті:
    і openpyxl, і xlrd дивляться на розширення й відмовляються.
    """
    import io as _io
    import openpyxl
    with open(шлях, 'rb') as f:
        байти = f.read()
    if байти[:2] != b'PK':
        return None
    try:
        wb = openpyxl.load_workbook(_io.BytesIO(байти), read_only=True, data_only=True)
    except Exception:
        return None
    ws = wb[wb.sheetnames[0]]
    рядки = list(ws.iter_rows(values_only=True))
    шапка_i = next((i for i, r in enumerate(рядки[:5])
                    if r and any('ID параметра' in str(c) for c in r if c)), None)
    if шапка_i is None:
        return None
    шапка = [str(c).strip() if c else '' for c in рядки[шапка_i]]
    i = {н: шапка.index(н) for н in шапка if н}
    пар = collections.OrderedDict()
    for row in рядки[шапка_i + 1:]:
        if not row or not row[i.get('ID параметра', 0)]:
            continue
        pid = str(row[i['ID параметра']]).strip()
        зап = пар.setdefault(pid, {
            'назва': str(row[i.get('Назва параметра', 1)] or '').strip(),
            'тип': str(row[i.get('Тип параметра', 2)] or '').strip(),
            'одиниця': str(row[i.get('Одиниця вимірювання', 4)] or '').strip(),
            'значення': []})
        зн = row[i['ID значення']] if 'ID значення' in i and i['ID значення'] < len(row) else None
        if зн is not None and str(зн).strip():
            зап['значення'].append(str(зн).strip())
    return пар


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--категорії')
    ap.add_argument('--з-каталогу-юк', action='store_true', dest='з_юк')
    a = ap.parse_args()
    cids = []
    if a.категорії:
        cids = [c.strip() for c in a.категорії.split(',') if c.strip()]
    if a.з_юк:
        d = json.load(open(os.path.join(BASE, 'data', 'yug_catalog.json'), encoding='utf-8'))
        лічильник = collections.Counter(
            str(t['rozetka_code']) for lst in (d.get('товари') or {}).values()
            for t in lst if t.get('rozetka_code'))
        cids = [c for c, _ in лічильник.most_common()]
    if not cids:
        sys.exit('не вказано категорій')
    h = {'Authorization': f'Bearer {RZ.get_token()}', 'Accept': 'application/json'}
    дані = json.load(open(OUT, encoding='utf-8')) if os.path.exists(OUT) else {}
    нових = 0
    for cid in cids:
        if cid in дані:
            continue
        шлях = завантажити(cid, h)
        if not шлях:
            print(f'{cid}: не віддалось', file=sys.stderr)
            continue
        розбір = розібрати(шлях)
        if розбір is None:
            print(f'{cid}: файл не читається', file=sys.stderr)
            continue
        дані[cid] = розбір
        нових += 1
        знач = sum(len(v['значення']) for v in розбір.values())
        print(f'{cid}: параметрів {len(розбір)}, допустимих значень {знач}')
        time.sleep(0.4)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(дані, f, ensure_ascii=False)
    print(f'\nкатегорій у довіднику: {len(дані)} (нових {нових}) → {OUT}')


if __name__ == '__main__':
    main()
