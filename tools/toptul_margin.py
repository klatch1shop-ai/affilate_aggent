#!/usr/bin/env python3
"""Юніт-економіка TOPTUL на Rozetka: де ми продаємо в мінус.

ФОРМУЛА (усі складники з фактів, не з припущень):

    закупівля  Z = РРЦ × (1 − знижка)        знижка TOPTUL = 12 %, з накладних
    комісія    C = ціна × ставка/100 × 1.08  ставка з /items-commissions/search
    доставка   D = 35 грн                     платить власник, Rozetka не компенсує
    маржа      M = ціна − Z − C − D
    ціна для маржі M:  P = (Z + D + M) / (1 − ставка×1.08/100)

ЧОМУ 1.08. `commission_sum` з API на всіх перевірених позиціях рівно на
8 % більший за `ціна × ставка`. Тобто оголошені 18 % коштують 19,44 %.
Що це за коефіцієнт — питання до Rozetka; у розрахунку беремо факт.

ЧОМУ ЛИШЕ TOPTUL. Знижка 12 % підтверджена на 4 накладних без жодного
відхилення. У PROTESTER на 4 накладних знижка 15–21,3 % — переносити її
на каталог не можна, доки не знаємо правила.

    venv/bin/python tools/toptul_margin.py --маржа 40
"""
import argparse
import glob
import json
import os
import re
import sys

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
КОМ = os.path.join(BASE, 'data', 'rozetka_commissions.json')
OUT = os.path.join(BASE, 'exports', 'toptul_margin.csv')
ДОСТАВКА = 35.0
КОЕФ = 1.08
ЗНИЖКА = {'TOPTUL': 0.12}


def прайс():
    """{артикул: (бренд, РРЦ)} з прайсу постачальника."""
    import openpyxl
    файли = glob.glob('/tmp/grand/Прайс 1*.xlsx') + glob.glob(
        os.path.join(BASE, 'data', 'toptul', 'private', 'Прайс*.xlsx'))
    if not файли:
        sys.exit('немає прайсу постачальника: поклади його в /tmp/grand/')
    wb = openpyxl.load_workbook(файли[0], read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out, бренд = {}, None
    for row in ws.iter_rows(min_row=4, values_only=True):
        арт, товар, ц = row[1], row[2], row[5]
        if not арт and товар and not ц:
            # У прайсі два рівні заголовків: бренд (TOPTUL, PROTESTER —
            # великими латиницею) і підрозділ категорії («Насадки, головки»).
            # Спершу я приймав за бренд будь-який заголовок, і бренд
            # затирався назвою категорії — у вибірку не потрапила жодна
            # позиція TOPTUL.
            т = str(товар).strip()
            if т.isupper() and re.fullmatch(r'[A-Z0-9 .&-]+', т):
                бренд = т
            continue
        if арт and ц:
            try:
                out[str(арт).strip()] = (бренд, float(str(ц).replace(',', '.')))
            except ValueError:
                pass
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--маржа', type=float, default=40.0)
    a = ap.parse_args()
    п = прайс()
    ком = json.load(open(КОМ, encoding='utf-8'))
    рядки, збиткові = [], 0
    for it in ком:
        арт = (it.get('item_article') or '').strip()
        дані = п.get(арт)
        if not дані:
            continue
        бренд, ррц = дані
        if бренд not in ЗНИЖКА:
            continue
        try:
            ціна = float(it['item_price'])
            ставка = float(it['commission'])
            сума_ком = float(it['commission_sum'])
        except (KeyError, ValueError):
            continue
        z = ррц * (1 - ЗНИЖКА[бренд])
        маржа = ціна - z - сума_ком - ДОСТАВКА
        дільник = 1 - ставка * КОЕФ / 100
        потрібна = (z + ДОСТАВКА + a.маржа) / дільник if дільник > 0 else 0
        рядки.append({'артикул': арт, 'бренд': бренд, 'назва': it.get('item_name', '')[:70],
                      'категорія': it.get('catalog_name'), 'ставка_%': ставка,
                      'ціна_зараз': round(ціна, 2), 'закупівля': round(z, 2),
                      'комісія': round(сума_ком, 2), 'доставка': ДОСТАВКА,
                      'маржа_зараз': round(маржа, 2),
                      f'ціна_для_{int(a.маржа)}': round(потрібна, 2),
                      'підняти_на_%': round((потрібна / ціна - 1) * 100, 1) if ціна else 0})
        збиткові += маржа < 0
    рядки.sort(key=lambda r: r['маржа_зараз'])
    import csv
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(рядки[0]))
        w.writeheader()
        w.writerows(рядки)
    прибуткові = [r for r in рядки if r['маржа_зараз'] >= a.маржа]
    print(f'позицій TOPTUL на Rozetka з відомою РРЦ: {len(рядки)}')
    print(f'  продаються у ЗБИТОК: {збиткові} ({збиткові * 100 // len(рядки)}%)')
    print(f'  маржа >= {a.маржа:.0f} грн: {len(прибуткові)} '
          f'({len(прибуткові) * 100 // len(рядки)}%)')
    print(f'  сумарна маржа зараз: {sum(r["маржа_зараз"] for r in рядки):,.0f} грн'.replace(',', ' '))
    print(f'\n{"артикул":12} {"ціна":>8} {"закуп":>8} {"комісія":>8} {"маржа":>8} {"треба":>8}')
    for r in рядки[:8]:
        print(f"{r['артикул'][:12]:12} {r['ціна_зараз']:8.2f} {r['закупівля']:8.2f} "
              f"{r['комісія']:8.2f} {r['маржа_зараз']:8.2f} {r[f'ціна_для_{int(a.маржа)}']:8.2f}")
    print(f'\n→ {OUT}')


if __name__ == '__main__':
    main()
