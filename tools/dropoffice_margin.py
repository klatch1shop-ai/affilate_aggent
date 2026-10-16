#!/usr/bin/env python3
"""Юніт-економіка dropoffice на Rozetka: що можна продавати без збитку.

Прайс dropoffice, на відміну від TOPTUL, містить і нашу закупівлю (PRICE),
і роздрібну (RETAIL PRICE) — тож нічого виводити не треба.

    маржа = ціна − закупівля − комісія − доставка
    комісія = ціна × ставка/100 × 1.08   (коефіцієнт перевірено на API)
    доставка = 35 грн, платить власник

    venv/bin/python tools/dropoffice_margin.py --прайс /шлях/price.xlsx --маржа 40
"""
import argparse
import csv
import json
import os

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
КОМ = os.path.join(BASE, 'data', 'rozetka_commissions.json')
OUT = os.path.join(BASE, 'exports', 'dropoffice_margin.csv')
ДОСТАВКА = 35.0
КОЕФ = 1.08


def число(s):
    if s is None:
        return 0.0
    s = str(s).replace('\xa0', '').replace(' ', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def прайс(шлях):
    wb = openpyxl.load_workbook(шлях, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    рядки = ws.iter_rows(values_only=True)
    шапка = [str(c).strip().upper() if c else '' for c in next(рядки)]
    i = {н: шапка.index(н) for н in ('ARTICLE', 'NAME', 'CATEGORY', 'PRICE',
                                      'RETAIL PRICE', 'QUANTITY') if н in шапка}
    out = {}
    for row in рядки:
        арт = str(row[i['ARTICLE']]).strip() if row[i['ARTICLE']] else ''
        if not арт:
            continue
        out[арт] = {'назва': str(row[i['NAME']] or '')[:70],
                    'категорія': str(row[i['CATEGORY']] or ''),
                    'закупівля': число(row[i['PRICE']]),
                    'ррц': число(row[i['RETAIL PRICE']]),
                    'наявність': str(row[i['QUANTITY']] or '')[:20]}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--прайс', required=True)
    ap.add_argument('--маржа', type=float, default=40.0)
    a = ap.parse_args()
    п = прайс(a.прайс)
    ком = {(x.get('item_article') or '').strip(): x
           for x in json.load(open(КОМ, encoding='utf-8'))}
    рядки = []
    for арт, д in п.items():
        it = ком.get(арт)
        if not it or not д['закупівля']:
            continue
        ціна = число(it['item_price'])
        ставка = число(it['commission'])
        сума = число(it['commission_sum'])
        маржа = ціна - д['закупівля'] - сума - ДОСТАВКА
        дільник = 1 - ставка * КОЕФ / 100
        потрібна = (д['закупівля'] + ДОСТАВКА + a.маржа) / дільник if дільник > 0 else 0
        рядки.append({'артикул': арт, 'назва': д['назва'], 'категорія': д['категорія'],
                      'наявність': д['наявність'], 'ставка_%': ставка,
                      'ціна_зараз': round(ціна, 2), 'закупівля': д['закупівля'],
                      'ррц_постачальника': д['ррц'], 'комісія': round(сума, 2),
                      'доставка': ДОСТАВКА, 'маржа_зараз': round(маржа, 2),
                      f'ціна_для_{int(a.маржа)}': round(потрібна, 2),
                      'підняти_%': round((потрібна / ціна - 1) * 100, 1) if ціна else 0})
    рядки.sort(key=lambda r: -r['маржа_зараз'])
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(рядки[0]))
        w.writeheader(); w.writerows(рядки)
    добрі = [r for r in рядки if r['маржа_зараз'] >= a.маржа]
    збиток = [r for r in рядки if r['маржа_зараз'] < 0]
    print(f'позицій dropoffice, що є і в прайсі, і на Rozetka: {len(рядки)} з {len(п)}')
    print(f'  маржа >= {a.маржа:.0f} грн: {len(добрі)} ({len(добрі)*100//len(рядки)}%)')
    print(f'  у збиток: {len(збиток)} ({len(збиток)*100//len(рядки)}%)')
    print(f'  сумарна маржа по одній штуці кожної: {sum(r["маржа_зараз"] for r in рядки):,.0f} грн'.replace(',', ' '))
    print(f'\n{"артикул":16} {"ціна":>7} {"закуп":>7} {"комісія":>8} {"маржа":>8}  назва')
    for r in рядки[:10]:
        print(f"{r['артикул'][:16]:16} {r['ціна_зараз']:7.0f} {r['закупівля']:7.0f} "
              f"{r['комісія']:8.2f} {r['маржа_зараз']:8.2f}  {r['назва'][:38]}")
    print(f'\n→ {OUT}')


if __name__ == '__main__':
    main()
